"""Human spot check of the judgments.

`sample` draws pairs without inherited labels and pairs with inherited labels,
mixes them, and writes a grading sheet that shows the same topic and metadata
text as the assessor batches, without any existing grade. `score` compares each
completed sheet (CSV or XLSX) with assessors A and B on all graded pairs and
with the inherited labels on the inherited pairs. With two or more sheets,
graded independently by different people, it also compares the sheets with
each other on the pairs graded in both.
"""
import argparse
import csv
import itertools
import json
import random
from pathlib import Path

import numpy as np

from build_judgment_pool import render
from analyse_judgments import agreement

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent


def sample(n_new, n_inherited, seed, output):
    key = json.loads((HERE / 'judgments/pool_key.json').read_text())['items']
    inherited = {q['id']: q['judgments'] for q in json.loads((ROOT / 'data/legacy/judgments.json').read_text())['queries']}
    topics = {t['id']: t for t in json.loads((HERE / 'judgments/topics.json').read_text())['topics']}
    catalog = json.loads((ROOT / 'data/legacy/catalog.json').read_text())
    new = sorted(i for i, v in key.items() if v['dataset_id'] not in inherited[v['query_id']])
    old = sorted(i for i, v in key.items() if v['dataset_id'] in inherited[v['query_id']])
    rng = random.Random(seed)
    chosen = rng.sample(new, n_new) + rng.sample(old, n_inherited)
    chosen.sort(key=lambda i: (key[i]['query_id'], i))
    rows = [['item', 'topic', 'query', 'information_need', 'metadata', 'grade (0/1/2)', 'note']]
    for i in chosen:
        t = topics[key[i]['query_id']]
        rows.append([i, t['id'], t['query'], t['description'], '\n'.join(render(catalog[key[i]['dataset_id']])), '', ''])
    if str(output).endswith('.xlsx'):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        wb = Workbook()
        ws = wb.active
        ws.title = 'grades'
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for column, width in zip('ABCDEFG', (11, 8, 30, 30, 90, 13, 30)):
            ws.column_dimensions[column].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
        wb.save(output)
    else:
        with open(output, 'w', newline='', encoding='utf-8-sig') as f:
            csv.writer(f).writerows(rows)
    print(f'{n_new} of {len(new)} new and {n_inherited} of {len(old)} inherited pairs written to {output}')


def read_sheet(sheet):
    if str(sheet).endswith('.xlsx'):
        from openpyxl import load_workbook
        values = list(load_workbook(sheet, read_only=True).active.values)
        header = [str(v) for v in values[0]]
        rows = [dict(zip(header, ['' if v is None else str(v) for v in r])) for r in values[1:]]
    else:
        with open(sheet, encoding='utf-8-sig') as f:
            rows = list(csv.DictReader(f))
    return {r['item']: int(float(r['grade (0/1/2)'])) for r in rows if r['grade (0/1/2)'].strip()}


def compare(human, key, inherited, assessors):
    items = sorted(human)
    old = [i for i in items if key[i]['dataset_id'] in inherited[key[i]['query_id']]]
    new = [i for i in items if i not in old]
    result = {'pairs': len(items), 'new_pairs': len(new), 'inherited_pairs': len(old)}
    for name, grades in assessors.items():
        result[f'vs_{name}_all'] = agreement([human[i] for i in items], [grades[i] for i in items])
        result[f'vs_{name}_new'] = agreement([human[i] for i in new], [grades[i] for i in new])
    result['vs_inherited'] = agreement([human[i] for i in old],
                                       [inherited[key[i]['query_id']][key[i]['dataset_id']] for i in old])
    return result


def score(sheets):
    key = json.loads((HERE / 'judgments/pool_key.json').read_text())['items']
    inherited = {q['id']: q['judgments'] for q in json.loads((ROOT / 'data/legacy/judgments.json').read_text())['queries']}
    assessors = {name: {j['item']: j['grade'] for j in json.loads((HERE / f'judgments/assessor_{name}.json').read_text())['judgments']}
                 for name in 'AB'}
    humans = {f'H{n}': read_sheet(sheet) for n, sheet in enumerate(sheets, 1)}
    result = {'sheets': {h: Path(sheet).name for h, sheet in zip(humans, sheets)},
              'humans': {h: compare(grades, key, inherited, assessors) for h, grades in humans.items()}}
    names = list(humans)
    for a, b in itertools.combinations(names, 2):
        common = sorted(set(humans[a]) & set(humans[b]))
        old = [i for i in common if key[i]['dataset_id'] in inherited[key[i]['query_id']]]
        new = [i for i in common if i not in old]
        result[f'{a}_vs_{b}'] = {
            'all': agreement([humans[a][i] for i in common], [humans[b][i] for i in common]),
            'new': agreement([humans[a][i] for i in new], [humans[b][i] for i in new]),
            'inherited': agreement([humans[a][i] for i in old], [humans[b][i] for i in old])}
    (HERE / 'results/spot_check.json').write_text(json.dumps(result, indent=2) + '\n')
    for h, block in result['humans'].items():
        for k, v in block.items():
            if isinstance(v, dict):
                print(f"{h} {k}: n={v['pairs']} exact={v['exact']:.3f} weighted kappa={v['weighted_kappa_quadratic']:.3f}")
    for k, v in result.items():
        if '_vs_' in k:
            for part, x in v.items():
                print(f"{k} {part}: n={x['pairs']} exact={x['exact']:.3f} weighted kappa={x['weighted_kappa_quadratic']:.3f}")


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest='command', required=True)
    s = sub.add_parser('sample')
    s.add_argument('--new', type=int, default=40)
    s.add_argument('--inherited', type=int, default=20)
    s.add_argument('--seed', type=int, default=2027)
    s.add_argument('--output', default=str(HERE / 'judgments/spot_check_sheet.xlsx'))
    c = sub.add_parser('score')
    c.add_argument('sheets', nargs='+', help='completed sheets, one per assessor')
    args = p.parse_args()
    if args.command == 'sample':
        sample(args.new, args.inherited, args.seed, args.output)
    else:
        score(args.sheets)


if __name__ == '__main__':
    main()
